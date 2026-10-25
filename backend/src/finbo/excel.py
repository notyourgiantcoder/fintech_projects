"""Excel export and import: the spreadsheet as a human-in-the-loop surface.

The exported workbook carries the user's numbers plus an Approvals sheet listing the
plans Finbo currently proposes, each with a Status dropdown. The user marks a plan
Approved or Sent back in Excel and imports the file; only proposal IDs the server itself
generated are accepted, so an edited file can't inject arbitrary decisions.
"""

from __future__ import annotations

import hashlib
import zipfile
from dataclasses import dataclass
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.datavalidation import DataValidation

from .core import Currency, format_amount
from .store import Approval, Workspace

STATUSES = ("Pending", "Approved", "Sent back")
HEADER = Font(bold=True)
HIGHLIGHT = PatternFill("solid", fgColor="FFD60A")


@dataclass(frozen=True)
class Proposal:
    id: str
    subject: str


def _id(subject: str) -> str:
    return hashlib.sha256(subject.encode()).hexdigest()[:12]


def debt_subject(debt: dict[str, Any], currency: Currency | str) -> str:
    """The one name for a debt plan, wherever it's proposed (app, chat, Excel, Telegram)."""
    method = debt["comparison"]["cheapest"]
    return (
        f"Debt plan: {method}, {format_amount(debt['monthly_budget'], Currency(currency))} a month"
    )


def proposals(calc: dict[str, Any]) -> list[Proposal]:
    currency = Currency(calc["currency"])
    out = []
    debt = calc["debt"]
    if "comparison" in debt:
        out.append(debt_subject(debt, currency))
    plan = calc["emergency_fund"]["plan"]
    if not plan["is_funded"] and plan["months_to_target"]:
        contribution = format_amount(plan["monthly_contribution"], currency)
        target = format_amount(plan["target_amount"], currency)
        out.append(
            f"Safety net: set aside {contribution} a month "
            f"to reach {target} in {plan['months_to_target']} months"
        )
    return [Proposal(id=_id(s), subject=s) for s in out]


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list[Any]], widths: list[int]):
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font = HEADER
    for row in rows:
        ws.append(row)
    for i, width in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = width
    ws.freeze_panes = "A2"
    return ws


def build_workbook(ws: Workspace, calc: dict[str, Any], decided: list[Approval]) -> bytes:
    currency = Currency(calc["currency"])
    spending = calc["spending"]
    plan = calc["emergency_fund"]["plan"]
    debt = calc["debt"]

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append(["Finbo plan export"])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append(["Educational guidance only, not regulated financial advice."])
    summary.append([])
    rows: list[tuple[str, Any]] = [
        ("Country", "India" if ws.market == "in" else "United States"),
        ("Currency", currency.value),
        ("Statement period", f"{spending['period_start']} to {spending['period_end']}"),
        ("Income a month", float(spending["monthly_income"])),
        ("Spending a month (average)", float(spending["avg_monthly_spending"])),
        ("Left over a month (average)", float(spending["avg_monthly_savings"])),
        ("Savings rate %", float(spending["savings_rate_pct"] or 0)),
        ("Emergency fund target", float(plan["target_amount"])),
        ("Emergency fund today", float(plan["current_balance"])),
        ("Emergency fund gap", float(plan["gap"])),
    ]
    if "comparison" in debt:
        best = debt["comparison"][debt["comparison"]["cheapest"]]
        rows += [
            ("Debt budget a month", float(debt["monthly_budget"])),
            ("Months to debt-free (cheapest method)", best["months"]),
            (
                "Interest saved with avalanche",
                float(debt["comparison"]["interest_saved_with_avalanche"]),
            ),
        ]
    for label, value in rows:
        summary.append([label, value])
    summary.column_dimensions["A"].width = 40
    summary.column_dimensions["B"].width = 22

    _sheet(
        wb,
        "Spending",
        ["Category", "Filed as", "Per month", "Share %"],
        [
            [c["category"], c["bucket"], float(c["monthly_avg"]), float(c["share_pct"])]
            for c in spending["categories"]
        ],
        [24, 16, 14, 10],
    )
    _sheet(
        wb,
        "Transactions",
        ["Date", "Category", "Amount", "Description"],
        [[t.date, t.category, float(t.amount), t.description] for t in ws.finance.transactions],
        [12, 20, 14, 36],
    )
    if "comparison" in debt:
        lanes = []
        for key in ("avalanche", "snowball"):
            p = debt["comparison"][key]
            for step in p["payoff_order"]:
                lanes.append(
                    [key, step["name"], step["month_paid_off"], float(step["interest_paid"])]
                )
        _sheet(
            wb,
            "Debt plan",
            ["Method", "Debt", "Paid off in month", "Interest paid"],
            lanes,
            [12, 24, 18, 16],
        )

    sheet = wb.create_sheet("Approvals")
    sheet.append(
        ["Change Status to Approved or Sent back, save the file, then import it in Finbo Settings."]
    )
    sheet["A1"].font = Font(italic=True)
    sheet.append(["ID", "Plan", "Status", "Note"])
    for cell in sheet[2]:
        cell.font = HEADER
    decided_subjects = {a.subject: a for a in decided}
    pending = [p for p in proposals(calc) if p.subject not in decided_subjects]
    for p in pending:
        sheet.append([p.id, p.subject, "Pending", ""])
    if pending:
        dv = DataValidation(type="list", formula1=f'"{",".join(STATUSES)}"', allow_blank=False)
        sheet.add_data_validation(dv)
        dv.add(f"C3:C{2 + len(pending)}")
    for a in decided:
        row = [
            "decided",
            a.subject,
            "Approved" if a.decision == "approved" else "Sent back",
            a.note,
        ]
        sheet.append(row)
        if a.decision == "approved":
            sheet.cell(row=sheet.max_row, column=3).fill = HIGHLIGHT
    sheet.column_dimensions["A"].width = 14
    sheet.column_dimensions["B"].width = 70
    sheet.column_dimensions["C"].width = 12
    sheet.column_dimensions["D"].width = 30

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


class WorkbookError(ValueError):
    pass


def read_decisions(data: bytes, current: list[Proposal]) -> list[tuple[Proposal, str, str]]:
    """Return (proposal, decision, note) for rows the user marked, matching live proposals."""
    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError) as e:
        raise WorkbookError("That file isn’t an Excel workbook (.xlsx).") from e
    if "Approvals" not in wb.sheetnames:
        raise WorkbookError(
            "The workbook has no Approvals sheet. Export a fresh copy and try again."
        )
    by_id = {p.id: p for p in current}
    out = []
    for row in wb["Approvals"].iter_rows(min_row=3, values_only=True):
        if not row or row[0] is None:
            continue
        pid, _plan, status, note = (list(row) + [None] * 4)[:4]
        proposal = by_id.get(str(pid).strip())
        status = str(status or "").strip().lower()
        if proposal is None or status not in {"approved", "sent back"}:
            continue
        decision = "approved" if status == "approved" else "sent_back"
        out.append((proposal, decision, str(note or "")[:1000]))
    return out
