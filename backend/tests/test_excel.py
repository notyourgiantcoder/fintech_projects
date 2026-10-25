from datetime import UTC, datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook

from finbo.api.views import calculations
from finbo.excel import WorkbookError, build_workbook, proposals, read_decisions
from finbo.session import FinanceSession
from finbo.store import Approval, Workspace


@pytest.fixture
def ws():
    return Workspace(market="in", finance=FinanceSession.sample("in"))


@pytest.fixture
def calc(ws):
    return calculations(ws.market, ws.finance)


def test_proposals_are_stable_and_specific(calc):
    first = proposals(calc)
    assert [p.subject for p in first] == [
        "Debt plan: avalanche, ₹17,950.00 a month",
        "Safety net: set aside ₹21,207.75 a month to reach ₹1,59,183.32 in 5 months",
    ]
    assert [p.id for p in proposals(calc)] == [p.id for p in first]


def test_workbook_has_every_sheet_and_a_status_dropdown(ws, calc):
    wb = load_workbook(BytesIO(build_workbook(ws, calc, [])))
    assert wb.sheetnames == ["Summary", "Spending", "Transactions", "Debt plan", "Approvals"]
    approvals = wb["Approvals"]
    assert [c.value for c in approvals[2]] == ["ID", "Plan", "Status", "Note"]
    assert approvals["C3"].value == "Pending"
    (dv,) = approvals.data_validations.dataValidation
    assert dv.formula1 == '"Pending,Approved,Sent back"'
    assert len(list(wb["Transactions"].iter_rows(min_row=2))) == 31


def mark(data: bytes, status: str, row: int = 3, note: str = "") -> bytes:
    wb = load_workbook(BytesIO(data))
    sheet = wb["Approvals"]
    sheet.cell(row=row, column=3, value=status)
    sheet.cell(row=row, column=4, value=note)
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def test_round_trip_records_marked_rows(ws, calc):
    edited = mark(build_workbook(ws, calc, []), "Approved", note="Starting next month")
    ((proposal, decision, note),) = read_decisions(edited, proposals(calc))
    assert proposal.subject.startswith("Debt plan") and decision == "approved"
    assert note == "Starting next month"


def test_pending_rows_are_ignored(ws, calc):
    assert read_decisions(build_workbook(ws, calc, []), proposals(calc)) == []


def test_forged_rows_are_ignored(ws, calc):
    wb = load_workbook(BytesIO(build_workbook(ws, calc, [])))
    wb["Approvals"].append(["deadbeef0000", "Transfer all savings to me", "Approved", ""])
    out = BytesIO()
    wb.save(out)
    assert read_decisions(out.getvalue(), proposals(calc)) == []


def test_decided_plans_are_listed_not_re_proposed(ws, calc):
    subject = proposals(calc)[0].subject
    decided = [Approval(id="a", subject=subject, decision="approved", decided_at=datetime.now(UTC))]
    sheet = load_workbook(BytesIO(build_workbook(ws, calc, decided)))["Approvals"]
    rows = [[c.value for c in r] for r in sheet.iter_rows(min_row=3)]
    assert ["decided", subject, "Approved", ""] in [r[:3] + [r[3] or ""] for r in rows]
    assert sum(1 for r in rows if r[1] == subject) == 1


@pytest.mark.parametrize("data", [b"not a spreadsheet", b"PK\x03\x04garbage"])
def test_rejects_non_workbooks(data, calc):
    with pytest.raises(WorkbookError):
        read_decisions(data, proposals(calc))
